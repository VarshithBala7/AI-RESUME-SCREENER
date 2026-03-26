import io
import re
import string
import uuid
from collections import Counter
from pathlib import Path
from typing import Iterable

from docx import Document
from fpdf import FPDF
from pypdf import PdfReader
STOPWORDS = {
    "the",
    "a",
    "an",
    "and",
    "or",
    "of",
    "to",
    "in",
    "for",
    "with",
    "on",
    "at",
    "by",
    "is",
    "are",
    "be",
    "as",
    "that",
    "this",
    "it",
    "your",
    "you",
    "from",
    "we",
    "our",
    "will",
    "have",
    "has",
    "using",
    "use",
    "job",
    "role",
    "experience",
    "years",
}
def parse_resume_bytes(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".txt":
        return data.decode("utf-8", errors="ignore")
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages)
    raise ValueError("Unsupported file format. Use PDF, DOCX, or TXT.")


def extract_keywords(text: str, limit: int = 40) -> list[str]:
    words = _tokenize(text)
    counter = Counter(words)
    return [word for word, _count in counter.most_common(limit)]


def calculate_match(job_description: str, resume_text: str) -> tuple[float, list[str], list[str]]:
    jd_keywords = set(extract_keywords(job_description, limit=60))
    resume_tokens = set(_tokenize(resume_text))

    if not jd_keywords:
        return 0.0, [], []

    matched = sorted(jd_keywords.intersection(resume_tokens))
    missing = sorted(jd_keywords.difference(resume_tokens))
    match_percentage = (len(matched) / len(jd_keywords)) * 100
    return round(match_percentage, 2), matched, missing


def generate_ats_resume_text(
    full_name: str,
    resume_text: str,
    job_description: str,
    missing_keywords: Iterable[str],
) -> str:
    jd_keywords = extract_keywords(job_description, limit=25)
    missing_keywords = list(missing_keywords)
    top_missing = missing_keywords[:15]

    sections = [
        f"{full_name}\nATS-Optimized Resume Draft",
        "PROFESSIONAL SUMMARY",
        (
            "Results-driven candidate with practical experience across technical and business workflows. "
            "Demonstrates adaptability, communication, and execution focus aligned to target role requirements."
        ),
        "CORE SKILLS",
        ", ".join(sorted(set(jd_keywords[:18]))),
        "EXPERIENCE HIGHLIGHTS",
        (
            "Delivered projects with measurable outcomes, collaborated in cross-functional teams, "
            "and used structured problem-solving to improve process quality."
        ),
        "PROJECTS",
        (
            "Implemented role-relevant solutions and documented technical decisions, testing methods, "
            "and production-ready recommendations."
        ),
        "EDUCATION",
        "Include your highest degree, institution, and graduation details here.",
    ]

    if top_missing:
        sections.extend(
            [
                "RECOMMENDED KEYWORDS TO INCORPORATE",
                ", ".join(top_missing),
            ]
        )

    sections.extend(
        [
            "SOURCE RESUME CONTEXT",
            _compact_text(resume_text, char_limit=1200),
            "TARGET JOB CONTEXT",
            _compact_text(job_description, char_limit=1200),
        ]
    )

    return "\n\n".join(sections)


def write_docx(content: str, output_path: Path) -> None:
    document = Document()
    for block in content.split("\n\n"):
        document.add_paragraph(block.strip())
    document.save(str(output_path))


def write_pdf(content: str, output_path: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    page_width = pdf.w - pdf.l_margin - pdf.r_margin
    safe_text = _sanitize_pdf_text(content).replace("\t", "    ")
    pdf.multi_cell(page_width, 7, safe_text if safe_text else " ")
    pdf.output(str(output_path))


def _tokenize(text: str) -> list[str]:
    cleaned = text.lower().translate(str.maketrans("", "", string.punctuation))
    tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9\-\+#]{1,}", cleaned)
    return [token for token in tokens if token not in STOPWORDS and len(token) > 2]


def _compact_text(text: str, char_limit: int) -> str:
    compact = re.sub(r"\s+", " ", text).strip()
    if len(compact) <= char_limit:
        return compact
    return compact[:char_limit] + "..."


def _sanitize_pdf_text(text: str) -> str:
    encoded = text.encode("latin-1", errors="ignore")
    return encoded.decode("latin-1")
